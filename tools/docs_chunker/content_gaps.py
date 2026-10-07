#!/usr/bin/env python
"""Close content gaps (offline, deterministic, read-only on C:\\HS-STEEL):
 - kind 'vba'          : VBA module source (oletools) + vba_procedures.jsonl
 - kind 'xlsx_formula' : formula text grouped by column pattern (R1C1-normalised) + named ranges
 - kind 'pdf_ocr'      : OCR (Windows.Media.Ocr ko) of PDF pages lacking a text layer
Merges into out/knowledge/chunks.jsonl (replacing earlier records of those kinds). Then run embed.py."""
import hashlib, json, re, subprocess, sys, tempfile
from pathlib import Path

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
import chunk as C

ROOT, OUT = C.ROOT, C.OUT
KINDS = ("vba", "xlsx_formula", "pdf_ocr")


def cid(kind, src, unit, i):
    return hashlib.sha256(f"{kind}|{src}|{unit}|{i}".encode("utf-8")).hexdigest()[:16]


def emit(recs, kind, src, sha, unit, text, page_or_sheet=None):
    for i, ch in enumerate(C.split(text)):
        recs.append({"id": cid(kind, src, unit, i), "source": src, "sha256": sha,
                     "page_or_sheet": page_or_sheet or unit, "kind": kind, "text": ch})


# ---------------- VBA ----------------
PROC = re.compile(r"^\s*(?:(Public|Private|Friend)\s+)?(?:Static\s+)?(Sub|Function|Property\s+(?:Get|Let|Set))\s+([^\s(]+)\s*(\(.*)?$", re.I)
END = re.compile(r"^\s*End\s+(Sub|Function|Property)\b", re.I)


def procedures(code):
    lines = code.replace("\r", "").split("\n")
    out, i = [], 0
    while i < len(lines):
        m = PROC.match(lines[i])
        if m and not lines[i].lstrip().startswith("'"):
            sig = lines[i].strip()
            start = i
            while sig.endswith("_") and i + 1 < len(lines):
                i += 1
                sig = sig[:-1].rstrip() + " " + lines[i].strip()
            j = i
            while j < len(lines) and not END.match(lines[j]):
                j += 1
            out.append((m.group(3), sig, j - start + 1))
            i = j
        i += 1
    return out


def do_vba(recs, procs, report):
    from oletools.olevba import VBA_Parser
    seen = {}
    files = sorted(p for p in ROOT.rglob("*.xlsm") if not p.name.startswith("~$"))
    for p in files:
        src = C.rel(p)
        sha = C.sha256_file(p)
        v = VBA_Parser(str(p))
        if not v.detect_vba_macros():
            v.close()
            continue
        report["vba_workbooks"] += 1
        mods = sorted(((n, code) for (_f, _s, n, code) in v.extract_macros()), key=lambda t: t[0])
        for name, code in mods:
            code = "\n".join(l for l in code.replace("\r", "").split("\n") if not l.startswith("Attribute VB_"))
            pr = procedures(code)
            for pn, sig, lc in pr:
                procs.append({"workbook": src, "module": name, "name": pn, "signature": sig, "line_count": lc})
            h = hashlib.sha256(code.encode("utf-8")).hexdigest()
            if h in seen:
                report["vba_duplicate_modules_skipped"] += 1
                continue
            seen[h] = (src, name)
            body = code.strip()
            if len(body) < 5:
                continue
            report["vba_modules"] += 1
            head = f"[VBA {src} :: module {name}] procedures: " + ", ".join(n for n, _, _ in pr)
            for i, ch in enumerate(C.split(body)):
                recs.append({"id": cid("vba", src, name, i), "source": src, "sha256": sha,
                             "page_or_sheet": name, "kind": "vba", "text": head + "\n" + ch})
        v.close()
    report["vba_procedures"] = len(procs)


# ---------------- formulas ----------------
from openpyxl.utils import column_index_from_string, get_column_letter

REF = re.compile(r"(?<![A-Za-z0-9_.!'])(\$?)([A-Z]{1,3})(\$?)(\d{1,7})(?![A-Za-z0-9_(])")


def r1c1(f, row, col):
    def sub(m):
        c = column_index_from_string(m.group(2))
        r = int(m.group(4))
        cs = f"C{c}" if m.group(1) else f"C[{c - col}]"
        rs = f"R{r}" if m.group(3) else f"R[{r - row}]"
        return rs + cs
    parts = re.split(r'("(?:[^"]|"")*")', f)
    return "".join(x if x.startswith('"') else REF.sub(sub, x) for x in parts)


def ftext(v):
    if hasattr(v, "text"):
        t = str(v.text)
        return t if t.startswith("=") else "=" + t
    return v


def runs(rows):
    out, s, prev = [], None, None
    for r in rows:
        if s is None:
            s = prev = r
        elif r == prev + 1:
            prev = r
        else:
            out.append((s, prev))
            s = prev = r
    if s is not None:
        out.append((s, prev))
    return ",".join(str(a) if a == b else f"{a}-{b}" for a, b in out)


def do_formulas(recs, report):
    import openpyxl
    files = sorted(p for p in ROOT.rglob("*") if p.suffix.lower() in (".xlsx", ".xlsm") and not p.name.startswith("~$"))
    for p in files:
        src = C.rel(p)
        sha = C.sha256_file(p)
        try:
            wb = openpyxl.load_workbook(str(p), data_only=False, keep_vba=False)
        except Exception as e:
            report["errors"].append({"source": src, "error": repr(e)})
            continue
        names = []
        try:
            for k, d in wb.defined_names.items():
                names.append(f"{k} = {d.attr_text}")
            for ws in wb.worksheets:
                for k, d in ws.defined_names.items():
                    names.append(f"{ws.title}!{k} = {d.attr_text}")
        except Exception:
            pass
        if names:
            report["named_ranges"] += len(names)
            emit(recs, "xlsx_formula", src, sha, "names", f"[{src}] Named ranges:\n" + "\n".join(sorted(names)), "[named ranges]")
        for ws in wb.worksheets:
            groups = {}
            for row in ws.iter_rows():
                for c in row:
                    f = ftext(c.value)
                    if isinstance(f, str) and f.startswith("="):
                        g = groups.setdefault((c.column, r1c1(f, c.row, c.column)), {"rows": [], "ex": (c.coordinate, f)})
                        g["rows"].append(c.row)
            if not groups:
                continue
            report["formula_cells"] += sum(len(g["rows"]) for g in groups.values())
            report["formula_patterns"] += len(groups)
            lines = [f"[{src}] Sheet '{ws.title}' formulas (grouped by column pattern; R1C1 normalised)"]
            for (col, norm), g in sorted(groups.items(), key=lambda kv: (kv[0][0], kv[1]["rows"][0], kv[0][1])):
                lines.append(f"Col {get_column_letter(col)}: {norm}  | e.g. {g['ex'][0]} {g['ex'][1]}  | rows {runs(g['rows'])} ({len(g['rows'])})")
            emit(recs, "xlsx_formula", src, sha, "f:" + ws.title, "\n".join(lines), ws.title)
        wb.close()


# ---------------- OCR ----------------
def do_ocr(recs, report):
    import pymupdf
    ps = sorted((ROOT / "HSSTEEL" / "도움말").glob("*.pdf"))
    tmp = Path(tempfile.mkdtemp(prefix="hs_ocr_"))
    todo = []
    for p in ps:
        src = C.rel(p)
        sha = C.sha256_file(p)
        d = pymupdf.open(str(p))
        for n, pg in enumerate(d, 1):
            if len(pg.get_text().strip()) < 40:
                key = hashlib.sha1(f"{src}|{n}".encode()).hexdigest()[:10]
                pg.get_pixmap(matrix=pymupdf.Matrix(3, 3)).save(str(tmp / f"{key}.png"))
                todo.append((src, sha, n, key))
    res = None
    out_json = tmp / "ocr.json"
    try:
        subprocess.run(["powershell", "-NoProfile", "-File", str(HERE / "ocr.ps1"), str(tmp), str(out_json), "ko"],
                       check=True, capture_output=True, timeout=900)
        res = json.loads(out_json.read_text(encoding="utf-8-sig"))
        if "error" in res:
            raise RuntimeError(res["error"])
        report["ocr_status"] = "ok (Windows.Media.Ocr ko)"
    except Exception as e:
        report["ocr_status"] = "ocr_unavailable"
        report["ocr_error"] = repr(e)
        res = None
    report["ocr_pages_attempted"] = len(todo)
    report["ocr_pages"] = []
    for src, sha, n, key in todo:
        txt = None if res is None else res.get(key, "")
        status = "ocr_unavailable" if txt is None else ("text" if txt.strip() else "blank_page")
        report["ocr_pages"].append({"source": src, "page": n, "status": status, "chars": len((txt or "").strip())})
        if txt and txt.strip():
            emit(recs, "pdf_ocr", src, sha, f"ocr:{n}", f"[OCR {src} p.{n}]\n{txt.strip()}", str(n))


def main():
    old = [json.loads(l) for l in open(OUT / "chunks.jsonl", encoding="utf-8")]
    keep = [r for r in old if r["kind"] not in KINDS]
    recs, procs = [], []
    report = {"vba_workbooks": 0, "vba_modules": 0, "vba_duplicate_modules_skipped": 0, "named_ranges": 0,
              "formula_cells": 0, "formula_patterns": 0, "errors": []}
    do_vba(recs, procs, report)
    do_formulas(recs, report)
    do_ocr(recs, report)
    allr = keep + recs
    allr.sort(key=lambda r: (r["source"], r["page_or_sheet"].zfill(6), r["id"]))
    with open(OUT / "chunks.jsonl", "w", encoding="utf-8", newline="\n") as f:
        for r in allr:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    procs.sort(key=lambda r: (r["workbook"], r["module"], r["name"], r["signature"]))
    with open(OUT / "vba_procedures.jsonl", "w", encoding="utf-8", newline="\n") as f:
        for r in procs:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    report["chunks_by_kind"] = {k: sum(1 for r in recs if r["kind"] == k) for k in KINDS}
    report["total_chunks"] = len(allr)
    (OUT / "content_gaps_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k: v for k, v in report.items() if k != "ocr_pages"}, ensure_ascii=False, indent=1))
    print(json.dumps(report.get("ocr_pages"), ensure_ascii=False))


if __name__ == "__main__":
    main()
